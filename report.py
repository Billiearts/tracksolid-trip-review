"""
report.py — trip detection, event summarisation, and Excel template filling.

Stop detection merges two sources:
  - Track points (jimi.device.track.list) — catches stops the parking API
    misses, e.g. deep-sleep overnight stops reported with a single sample.
  - Parking events (jimi.open.platform.report.parking) — accurate boundaries
    and an address when it fires.

Both sources filter their own stops at >= MIN_STOP_SECONDS before merging.
Track detection uses raw run durations only (no extension), which prevents
the historical over-extension bug where a 3-4 minute stop was inflated past
the 5-minute threshold.
"""

from __future__ import annotations

import json
import logging
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from jimi import (
    api_call, fetch_alarms, fetch_alarms_chunked,
    fetch_parking, RateLimitExceeded,
)


log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

GBC_GEOID = "791a3bb604ef41cfa302e5dc8f2c356c"
GBC_FENCE_NAME_FALLBACK = "TRIP TIME GBC"

SPEED_LIMIT_KPH = 75.0

ALERT_ACCELERATION = "41"
ALERT_CORNERING    = "43"
ALERT_DECELERATION = "48"
ALERT_OVERSPEED    = "overSpeed"

MIN_STOP_SECONDS     = 300   # stops >= 5 min count toward the report
STATIONARY_SPEED_KPH = 1.5
MERGE_GAP_SECONDS    = 10

FMT = "%Y-%m-%d %H:%M:%S"

# Excel number format for durations — [h] allows hours > 24.
DURATION_FORMAT = "[h]:mm:ss"


# ---------------------------------------------------------------------------
# Time helpers
# ---------------------------------------------------------------------------

def utc_str(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime(FMT)


def parse_utc(s: str) -> datetime:
    return datetime.strptime(s, FMT).replace(tzinfo=timezone.utc)


def _to_float(x) -> float:
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def _speed_of(p: dict) -> float:
    for key in ("gpsSpeed", "speed"):
        v = p.get(key)
        if v is not None:
            try:
                return float(v)
            except (TypeError, ValueError):
                pass
    return 0.0


def _is_stationary_point(p: dict) -> bool:
    ignition_off = (p.get("ignition") or "").upper() == "OFF"
    slow = _speed_of(p) < STATIONARY_SPEED_KPH
    return ignition_off and slow


def history_window(reference: datetime) -> tuple[datetime, datetime]:
    first_of_this_month = reference.replace(
        day=1, hour=0, minute=0, second=0, microsecond=0)
    end = reference
    y, m = first_of_this_month.year, first_of_this_month.month - 3
    while m <= 0:
        m += 12
        y -= 1
    calendar_begin = first_of_this_month.replace(year=y, month=m)
    retention_floor = reference - timedelta(days=89)
    begin = max(calendar_begin, retention_floor)
    return begin, end


# ---------------------------------------------------------------------------
# Trip detection (GBC geofence alarms)
# ---------------------------------------------------------------------------

def _is_gbc_alarm(alarm: dict) -> bool:
    if (alarm.get("geoid") or "").strip() == GBC_GEOID:
        return True
    return GBC_FENCE_NAME_FALLBACK in (alarm.get("alarmTypeName") or "")


def _gbc_alarms_in_window(alarms: list[dict],
                          begin: datetime, end: datetime) -> list[dict]:
    gbc = [
        a for a in alarms
        if _is_gbc_alarm(a) and begin <= parse_utc(a["alertTime"]) <= end
    ]
    gbc.sort(key=lambda a: a["alertTime"])
    return gbc


def _trips_from_gbc_alarms(gbc: list[dict]) -> list[dict]:
    trips: list[dict] = []
    open_start: datetime | None = None
    for a in gbc:
        event = (a.get("alertTypeId") or "").lower()
        t = parse_utc(a["alertTime"])
        if event == "in":
            if open_start is not None:
                log.warning("unclosed trip starting %s; restarting", open_start)
            open_start = t
        elif event == "out":
            if open_start is None:
                continue
            trips.append({
                "start": open_start, "end": t,
                "duration_s": (t - open_start).total_seconds(),
            })
            open_start = None
    return trips


def find_trips(imei: str, begin: datetime, end: datetime) -> list[dict]:
    alarms = fetch_alarms_chunked(imei, begin, end)
    gbc = _gbc_alarms_in_window(alarms, begin, end)
    return _trips_from_gbc_alarms(gbc)


# ---------------------------------------------------------------------------
# Track fetch (was originally defined here, not in jimi.py)
# ---------------------------------------------------------------------------

def fetch_track(imei: str, begin: datetime, end: datetime) -> list[dict]:
    response = api_call("jimi.device.track.list", {
        "imei":       imei,
        "begin_time": utc_str(begin),
        "end_time":   utc_str(end),
    })
    result = response.get("result") or response.get("data") or []
    if isinstance(result, dict):
        result = result.get("list", [])
    return result or []


# ---------------------------------------------------------------------------
# Stop detection — track source
# ---------------------------------------------------------------------------

def _track_stop_intervals(track: list[dict],
                          trip_start: datetime | None,
                          trip_end: datetime | None,
                          debug: bool = False) -> list[dict]:
    """
    Detect stops from raw track points. Uses raw run durations only —
    no extension, no trim. This means a stop whose raw duration is under
    MIN_STOP_SECONDS is dropped, which is correct: we don't inflate short
    stops to look longer than they are.
    """
    if not track:
        return []

    points: list[tuple[datetime, bool]] = []
    for p in track:
        ts = p.get("gpsTime")
        if not ts:
            continue
        try:
            dt = parse_utc(ts)
        except ValueError:
            continue
        points.append((dt, _is_stationary_point(p)))
    points.sort(key=lambda x: x[0])

    if trip_start:
        points = [p for p in points if p[0] >= trip_start]
    if trip_end:
        points = [p for p in points if p[0] <= trip_end]
    if not points:
        return []

    if debug:
        n_stat = sum(1 for _, s in points if s)
        log.debug("track points (filtered): %d  stationary=%d",
                  len(points), n_stat)

    raw_runs: list[tuple[int, int]] = []
    i = 0
    while i < len(points):
        if points[i][1]:
            j = i
            while j + 1 < len(points) and points[j + 1][1]:
                j += 1
            raw_runs.append((i, j))
            i = j + 1
        else:
            i += 1

    if debug:
        log.debug("stationary runs: %d", len(raw_runs))
        for i0, j0 in raw_runs:
            t0, t1 = points[i0][0], points[j0][0]
            log.debug("    %s → %s  (%.0fs, %d pt(s))",
                      t0.strftime("%Y-%m-%d %H:%M:%S"),
                      t1.strftime("%Y-%m-%d %H:%M:%S"),
                      (t1 - t0).total_seconds(), j0 - i0 + 1)

    intervals: list[dict] = []
    for i0, j0 in raw_runs:
        start = points[i0][0]
        end = points[j0][0]
        dur = (end - start).total_seconds()
        if dur < MIN_STOP_SECONDS:
            if debug:
                log.debug("DROP %s → %s  (%.0fs): under %ds",
                          start.strftime("%Y-%m-%d %H:%M:%S"),
                          end.strftime("%Y-%m-%d %H:%M:%S"),
                          dur, MIN_STOP_SECONDS)
            continue
        intervals.append({
            "start": start,
            "end": end,
            "seconds": dur,
            "address": "",
        })
    return intervals


# ---------------------------------------------------------------------------
# Stop detection — parking source
# ---------------------------------------------------------------------------

def _parking_stop_intervals(parking_events: list[dict],
                            trip_start: datetime | None,
                            trip_end: datetime | None) -> list[dict]:
    out: list[dict] = []
    for e in parking_events:
        s = e.get("start")
        en = e.get("end")
        if s is None or en is None:
            continue
        if trip_start and s < trip_start:
            continue
        if trip_end and s > trip_end:
            continue
        dur = e.get("seconds", 0)
        if dur < MIN_STOP_SECONDS:
            continue
        out.append({
            "start": s,
            "end": en,
            "seconds": dur,
            "address": e.get("address", "") or "",
        })
    return out


# ---------------------------------------------------------------------------
# Merge stop intervals from both sources
# ---------------------------------------------------------------------------

def _merge_stop_intervals(parking_intervals: list[dict],
                          track_intervals: list[dict]) -> list[dict]:
    """
    Merge overlapping intervals. Within a merged group:
      - Times: earliest start → latest end (union).
      - Address: the parking address if any source in the group had one,
                 otherwise blank.
    """
    items: list[dict] = []
    for p in parking_intervals:
        items.append({
            "start": p["start"], "end": p["end"],
            "address": p.get("address", ""), "src": "p",
        })
    for t in track_intervals:
        items.append({
            "start": t["start"], "end": t["end"],
            "address": t.get("address", ""), "src": "t",
        })

    if not items:
        return []

    items.sort(key=lambda x: x["start"])

    groups: list[list[dict]] = [[items[0]]]
    for item in items[1:]:
        group = groups[-1]
        group_end = max(x["end"] for x in group)
        if item["start"] <= group_end:
            group.append(item)
        else:
            groups.append([item])

    out: list[dict] = []
    for group in groups:
        start = min(x["start"] for x in group)
        end = max(x["end"] for x in group)
        address = ""
        for x in group:
            if x["src"] == "p" and x["address"]:
                address = x["address"]
                break
        out.append({
            "start": start,
            "end": end,
            "seconds": (end - start).total_seconds(),
            "address": address,
        })
    return out


# ---------------------------------------------------------------------------
# Public API: summarise stops from both sources
# ---------------------------------------------------------------------------

def summarise_stops(track: list[dict],
                    parking_events: list[dict],
                    trip_start: datetime | None = None,
                    trip_end: datetime | None = None,
                    debug: bool = False) -> dict:
    """
    Return stop summary, merging track-based and parking-based intervals.

    Return shape matches the previous version — all keys preserved.
    """
    track_intervals = _track_stop_intervals(
        track, trip_start, trip_end, debug=debug)
    parking_intervals = _parking_stop_intervals(
        parking_events, trip_start, trip_end)

    if debug:
        log.debug("track intervals:   %d", len(track_intervals))
        log.debug("parking intervals: %d", len(parking_intervals))

    merged = _merge_stop_intervals(parking_intervals, track_intervals)
    merged = [x for x in merged if x["seconds"] >= MIN_STOP_SECONDS]
    merged.sort(key=lambda x: x["start"])

    if debug:
        log.debug("merged stops: %d", len(merged))
        for s in merged:
            log.debug("    %s → %s  (%.0fs)",
                      s["start"].strftime("%Y-%m-%d %H:%M:%S"),
                      s["end"].strftime("%Y-%m-%d %H:%M:%S"),
                      s["seconds"])

    durations = [x["seconds"] for x in merged]

    return {
        "count":                 len(merged),
        "total_seconds":         sum(durations),
        "longest_seconds":       max(durations) if durations else 0,
        "durations":             list(durations),
        "intervals":             merged,
        # Mandatory rest break — logic removed for now.
        "rest_break_count":      0,
        "rest_break_seconds":    0,
        "rest_break_longest_s":  0,
        "rest_durations":        [],
        "rest_intervals":        [],
    }


def _empty_stops() -> dict:
    return {
        "count": 0, "total_seconds": 0, "longest_seconds": 0,
        "durations": [], "intervals": [],
        "rest_break_count": 0, "rest_break_seconds": 0,
        "rest_break_longest_s": 0,
        "rest_durations": [], "rest_intervals": [],
    }


# ---------------------------------------------------------------------------
# Event summarisation (unchanged — based on alarms)
# ---------------------------------------------------------------------------

def summarise_events(alarms: list[dict]) -> dict:
    counts = {
        "hard_braking":     0,
        "hard_cornering":   0,
        "speeding_10":      0,
        "speeding_20":      0,
        "speeding_over_20": 0,
    }
    max_speed = {k: 0.0 for k in counts}
    acceleration_count = 0
    event_list: list[dict] = []

    for a in alarms:
        atid = (a.get("alertTypeId") or "").strip()
        name = (a.get("alarmTypeName") or "").lower()
        speed = _to_float(a.get("speed"))
        t = a.get("alertTime")

        if atid == ALERT_DECELERATION or "sudden deceleration" in name:
            counts["hard_braking"] += 1
            max_speed["hard_braking"] = max(max_speed["hard_braking"], speed)
            event_list.append({"time": t, "type": "Hard braking", "speed": speed})
        elif atid == ALERT_CORNERING or "sharp turn" in name or "cornering" in name:
            counts["hard_cornering"] += 1
            max_speed["hard_cornering"] = max(max_speed["hard_cornering"], speed)
            event_list.append({"time": t, "type": "Hard cornering", "speed": speed})
        elif atid == ALERT_ACCELERATION or "sudden acceleration" in name:
            acceleration_count += 1
            event_list.append({"time": t, "type": "Sudden acceleration", "speed": speed})
        elif atid == ALERT_OVERSPEED or "overspeed" in name:
            over = speed - SPEED_LIMIT_KPH
            if over <= 0:
                continue
            bucket = ("speeding_10" if over < 10 else
                      "speeding_20" if over < 20 else
                      "speeding_over_20")
            counts[bucket] += 1
            max_speed[bucket] = max(max_speed[bucket], speed)
            label = {
                "speeding_10":      "Speeding <10 kph over",
                "speeding_20":      "Speeding <20 kph over",
                "speeding_over_20": "Speeding >20 kph over",
            }[bucket]
            event_list.append({"time": t, "type": label, "speed": speed})

    event_list.sort(key=lambda e: e.get("time") or "")

    return {
        "counts": counts,
        "max_speed": max_speed,
        "acceleration_count": acceleration_count,
        "events": event_list,
    }


# ---------------------------------------------------------------------------
# History
# ---------------------------------------------------------------------------

def build_history(imei: str, reference: datetime, *,
                  alarms: list[dict] | None = None,
                  hist_begin: datetime | None = None,
                  hist_end: datetime | None = None) -> dict | None:
    if hist_begin is None or hist_end is None:
        hist_begin, hist_end = history_window(reference)
    log.info("history window %s → %s", utc_str(hist_begin), utc_str(hist_end))

    if alarms is None:
        try:
            alarms = fetch_alarms_chunked(imei, hist_begin, hist_end)
        except RateLimitExceeded as exc:
            log.warning("history skipped — rate limit hit: %s", exc)
            return None

    gbc = _gbc_alarms_in_window(alarms, hist_begin, hist_end)
    trips = _trips_from_gbc_alarms(gbc)
    trips = [t for t in trips if hist_begin <= t["start"] < hist_end]
    log.info("%d trip(s) starting in window", len(trips))

    stop_counts:   list[int]   = []
    stop_totals_s: list[float] = []

    behav_events = {k: 0 for k in
                    ("hard_braking", "hard_cornering",
                     "speeding_10", "speeding_20", "speeding_over_20")}
    behav_max = {k: 0.0 for k in behav_events}
    behav_per_trip_max = {k: [] for k in behav_events}

    for t in trips:
        trip_alarms = [
            a for a in alarms
            if t["start"] <= parse_utc(a["alertTime"]) <= t["end"]
        ]

        try:
            parking = fetch_parking(imei, t["start"], t["end"])
        except RateLimitExceeded:
            log.warning("history: parking skipped for trip %s (rate limit)",
                        t["start"])
            parking = []
        except Exception as exc:
            log.warning("history: parking fetch failed for %s: %s",
                        t["start"], exc)
            parking = []

        try:
            track = fetch_track(imei, t["start"], t["end"])
        except RateLimitExceeded:
            log.warning("history: track skipped for trip %s (rate limit)",
                        t["start"])
            track = []
        except Exception as exc:
            log.warning("history: track fetch failed for %s: %s",
                        t["start"], exc)
            track = []

        s = summarise_stops(track, parking, t["start"], t["end"], debug=False)
        e = summarise_events(trip_alarms)

        stop_counts.append(s["count"])
        stop_totals_s.append(s["total_seconds"])

        for k in behav_events:
            behav_events[k] += e["counts"][k]
            behav_max[k] = max(behav_max[k], e["max_speed"][k])
            if e["counts"][k] > 0:
                behav_per_trip_max[k].append(e["max_speed"][k])

    def _avg(xs):
        return sum(xs) / len(xs) if xs else 0.0

    return {
        "trip_count": len(trips),
        "window_begin": hist_begin,
        "window_end":   hist_end,
        "stops": {
            "avg_count":    _avg(stop_counts),
            "max_count":    max(stop_counts) if stop_counts else 0,
            "min_count":    min(stop_counts) if stop_counts else 0,
            "avg_seconds":  _avg(stop_totals_s),
            "max_seconds":  max(stop_totals_s) if stop_totals_s else 0.0,
            "min_seconds":  min(stop_totals_s) if stop_totals_s else 0.0,
            "longest_single_stop_s":  0.0,
            "shortest_single_stop_s": 0.0,
        },
        "rests": {
            "avg_seconds": 0.0, "max_seconds": 0.0, "min_seconds": 0.0,
            "longest_single_rest_s": 0.0, "shortest_single_rest_s": 0.0,
        },
        "behaviours": {
            k: {
                "total_events":  behav_events[k],
                "max_speed":     behav_max[k],
                "avg_max_speed": _avg(behav_per_trip_max[k]),
            }
            for k in behav_events
        },
    }


# ---------------------------------------------------------------------------
# Excel filling
# ---------------------------------------------------------------------------

CELL_REPORT_GENERATED = "D4"
CELL_TRIP_START       = "D5"
CELL_TRIP_END         = "D6"
CELL_TRUCK_ID         = "D7"
CELL_DRIVER_NAME      = "D8"

CELL_REST_BREAK_HOURS = "E12"
CELL_STOPS_COUNT      = "E13"
CELL_STOPS_HOURS      = "E14"

CELL_HARD_BRAKING_COUNT = "E18"
CELL_HARD_BRAKING_SPEED = "F18"
CELL_HARD_CORNER_COUNT  = "E19"
CELL_HARD_CORNER_SPEED  = "F19"
CELL_SPEED_10_COUNT     = "E20"
CELL_SPEED_10_SPEED     = "F20"
CELL_SPEED_20_COUNT     = "E21"
CELL_SPEED_20_SPEED     = "F21"
CELL_SPEED_OVER_COUNT   = "E22"
CELL_SPEED_OVER_SPEED   = "F22"

CELL_HIST_REST_AVG_HOURS   = "H12"
CELL_HIST_REST_LONGEST_H   = "I12"
CELL_HIST_REST_SHORTEST_H  = "J12"

CELL_HIST_STOPS_AVG        = "H13"
CELL_HIST_STOPS_LONGEST    = "I13"
CELL_HIST_STOPS_SHORTEST   = "J13"

CELL_HIST_STOPS_AVG_HOURS     = "H14"
CELL_HIST_STOPS_LONGEST_HOURS = "I14"
CELL_HIST_STOPS_SHORTEST_HOURS = "J14"

CELL_HIST_HB_TOTAL = "H18"; CELL_HIST_HB_MAX = "I18"; CELL_HIST_HB_AVG = "J18"
CELL_HIST_HC_TOTAL = "H19"; CELL_HIST_HC_MAX = "I19"; CELL_HIST_HC_AVG = "J19"
CELL_HIST_S10_TOTAL = "H20"; CELL_HIST_S10_MAX = "I20"; CELL_HIST_S10_AVG = "J20"
CELL_HIST_S20_TOTAL = "H21"; CELL_HIST_S20_MAX = "I21"; CELL_HIST_S20_AVG = "J21"
CELL_HIST_SOV_TOTAL = "H22"; CELL_HIST_SOV_MAX = "I22"; CELL_HIST_SOV_AVG = "J22"


def _safe_write(ws, coord: str, value, number_format: str | None = None) -> None:
    target = coord
    for mr in ws.merged_cells.ranges:
        if coord in mr:
            anchor = str(mr).split(":")[0]
            if anchor != coord:
                log.debug("%s inside merge %s; → %s", coord, mr, anchor)
                target = anchor
            break
    ws[target] = value
    if number_format:
        ws[target].number_format = number_format


def _duration(seconds) -> float:
    try:
        s = float(seconds or 0)
    except (TypeError, ValueError):
        s = 0.0
    return s / 86400.0


def fill_template(template_path: Path, output_path: Path, *,
                  truck_id: str, driver_name: str,
                  trip_start: datetime | None,
                  trip_end: datetime | None,
                  stops: dict, events: dict,
                  history: dict | None) -> None:
    wb = load_workbook(template_path)
    ws = wb["Sheet1"]

    _safe_write(ws, CELL_REPORT_GENERATED,
                datetime.now(timezone.utc).strftime(FMT))
    _safe_write(ws, CELL_TRUCK_ID,    truck_id)
    _safe_write(ws, CELL_DRIVER_NAME, driver_name)
    if trip_start:
        _safe_write(ws, CELL_TRIP_START, trip_start.strftime(FMT))
    if trip_end:
        _safe_write(ws, CELL_TRIP_END,   trip_end.strftime(FMT))

    _safe_write(ws, CELL_REST_BREAK_HOURS,
                _duration(stops["rest_break_seconds"]), DURATION_FORMAT)
    _safe_write(ws, CELL_STOPS_COUNT, stops["count"])
    _safe_write(ws, CELL_STOPS_HOURS,
                _duration(stops["total_seconds"]), DURATION_FORMAT)

    c, m = events["counts"], events["max_speed"]
    _safe_write(ws, CELL_HARD_BRAKING_COUNT, c["hard_braking"])
    _safe_write(ws, CELL_HARD_BRAKING_SPEED, round(m["hard_braking"], 1))
    _safe_write(ws, CELL_HARD_CORNER_COUNT,  c["hard_cornering"])
    _safe_write(ws, CELL_HARD_CORNER_SPEED,  round(m["hard_cornering"], 1))
    _safe_write(ws, CELL_SPEED_10_COUNT,     c["speeding_10"])
    _safe_write(ws, CELL_SPEED_10_SPEED,     round(m["speeding_10"], 1))
    _safe_write(ws, CELL_SPEED_20_COUNT,     c["speeding_20"])
    _safe_write(ws, CELL_SPEED_20_SPEED,     round(m["speeding_20"], 1))
    _safe_write(ws, CELL_SPEED_OVER_COUNT,   c["speeding_over_20"])
    _safe_write(ws, CELL_SPEED_OVER_SPEED,   round(m["speeding_over_20"], 1))

    if history and history["trip_count"] > 0:
        h = history
        _safe_write(ws, CELL_HIST_REST_AVG_HOURS,
                    _duration(h["rests"]["avg_seconds"]), DURATION_FORMAT)
        _safe_write(ws, CELL_HIST_REST_LONGEST_H,
                    _duration(h["rests"]["max_seconds"]), DURATION_FORMAT)
        _safe_write(ws, CELL_HIST_REST_SHORTEST_H,
                    _duration(h["rests"]["min_seconds"]), DURATION_FORMAT)

        _safe_write(ws, CELL_HIST_STOPS_AVG,      round(h["stops"]["avg_count"], 1))
        _safe_write(ws, CELL_HIST_STOPS_LONGEST,  h["stops"]["max_count"])
        _safe_write(ws, CELL_HIST_STOPS_SHORTEST, h["stops"]["min_count"])
        _safe_write(ws, CELL_HIST_STOPS_AVG_HOURS,
                    _duration(h["stops"]["avg_seconds"]), DURATION_FORMAT)
        _safe_write(ws, CELL_HIST_STOPS_LONGEST_HOURS,
                    _duration(h["stops"]["max_seconds"]), DURATION_FORMAT)
        _safe_write(ws, CELL_HIST_STOPS_SHORTEST_HOURS,
                    _duration(h["stops"]["min_seconds"]), DURATION_FORMAT)

        b = h["behaviours"]
        _safe_write(ws, CELL_HIST_HB_TOTAL, b["hard_braking"]["total_events"])
        _safe_write(ws, CELL_HIST_HB_MAX,   round(b["hard_braking"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_HB_AVG,   round(b["hard_braking"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_HC_TOTAL, b["hard_cornering"]["total_events"])
        _safe_write(ws, CELL_HIST_HC_MAX,   round(b["hard_cornering"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_HC_AVG,   round(b["hard_cornering"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_S10_TOTAL, b["speeding_10"]["total_events"])
        _safe_write(ws, CELL_HIST_S10_MAX,   round(b["speeding_10"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_S10_AVG,   round(b["speeding_10"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_S20_TOTAL, b["speeding_20"]["total_events"])
        _safe_write(ws, CELL_HIST_S20_MAX,   round(b["speeding_20"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_S20_AVG,   round(b["speeding_20"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_SOV_TOTAL, b["speeding_over_20"]["total_events"])
        _safe_write(ws, CELL_HIST_SOV_MAX,   round(b["speeding_over_20"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_SOV_AVG,   round(b["speeding_over_20"]["avg_max_speed"], 1))

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Top-level pipeline
# ---------------------------------------------------------------------------

def compute_report(*, imei: str, truck_id: str, driver_name: str,
                   begin: datetime, end: datetime,
                   include_history: bool = True) -> dict:
    if include_history:
        hist_begin, hist_end = history_window(begin)
    else:
        hist_begin = hist_end = None

    union_begin = min(begin, hist_begin) if hist_begin else begin
    union_end   = max(end,   hist_end)   if hist_end   else end

    log.info("[1/6] Fetching alarms for %s → %s ...",
             utc_str(union_begin), utc_str(union_end))
    try:
        all_alarms = fetch_alarms_chunked(imei, union_begin, union_end)
    except RateLimitExceeded:
        raise
    log.info("      %d alarms fetched", len(all_alarms))

    gbc = _gbc_alarms_in_window(all_alarms, begin, end)
    trips = _trips_from_gbc_alarms(gbc)
    log.info("      %d trip(s) in search window", len(trips))

    if trips:
        trip = trips[0]
        trip_start, trip_end = trip["start"], trip["end"]
    else:
        trip_start = trip_end = None

    log.info("[2/6] Fetching track ...")
    try:
        track = fetch_track(imei, trip_start or begin,
                            trip_end or end) if trip_start else []
    except RateLimitExceeded:
        track = []
    log.info("      %d track points", len(track))

    log.info("[3/6] Fetching parking events ...")
    try:
        parking = fetch_parking(imei, trip_start or begin,
                                trip_end or end) if trip_start else []
    except RateLimitExceeded:
        parking = []
    log.info("      %d parking events", len(parking))

    log.info("[4/6] Filtering alarms for trip window ...")
    if trip_start and trip_end:
        alarms = [
            a for a in all_alarms
            if trip_start <= parse_utc(a["alertTime"]) <= trip_end
        ]
    else:
        alarms = []
    log.info("      %d alarms in trip", len(alarms))

    log.info("[5/6] Summarising ...")
    stops  = summarise_stops(track, parking, trip_start, trip_end, debug=False)
    events = summarise_events(alarms)
    log.info("      stops >= %dm: %d  total %.2fh",
             MIN_STOP_SECONDS // 60, stops["count"],
             stops["total_seconds"] / 3600)

    history = None
    if include_history and trip_start:
        log.info("[6/6] Building 3-month history (reusing fetched alarms) ...")
        try:
            history = build_history(
                imei, trip_start,
                alarms=all_alarms,
                hist_begin=hist_begin,
                hist_end=hist_end,
            )
        except RateLimitExceeded as exc:
            log.warning("history skipped (rate limit): %s", exc)
            history = None
        except Exception as exc:
            log.warning("history failed: %s", exc)

    return {
        "truck_id":      truck_id,
        "driver_name":   driver_name,
        "imei":          imei,
        "query_begin":   begin,
        "query_end":     end,
        "trip_start":    trip_start,
        "trip_end":      trip_end,
        "stops":         stops,
        "events":        events,
        "history":       history,
        "trips":         trips,
        "alarm_count":   len(alarms),
        "parking_count": len(parking),
        "track_count":   len(track),
        "generated_at":  datetime.now(timezone.utc),
    }


def write_excel(computed: dict, template_path: Path, output_path: Path,
                *,
                created_by: str | None = None,
                created_by_email: str | None = None) -> None:
    log.info("[write] Writing %s ...", output_path)
    fill_template(
        template_path, output_path,
        truck_id=computed["truck_id"],
        driver_name=computed["driver_name"],
        trip_start=computed["trip_start"],
        trip_end=computed["trip_end"],
        stops=computed["stops"],
        events=computed["events"],
        history=computed["history"],
    )

    ts = computed["trip_start"]
    te = computed["trip_end"]
    stops = computed["stops"]
    history = computed["history"]
    meta = {
        "truck_id":         computed["truck_id"],
        "driver_name":      computed["driver_name"],
        "imei":             computed["imei"],
        "trip_start":       utc_str(ts) if ts else None,
        "trip_end":         utc_str(te) if te else None,
        "query_begin":      utc_str(computed["query_begin"]),
        "query_end":        utc_str(computed["query_end"]),
        "generated_at":     utc_str(computed["generated_at"]),
        "trips_found":      len(computed["trips"]),
        "alarm_count":      computed["alarm_count"],
        "parking_count":    computed.get("parking_count", 0),
        "track_count":      computed.get("track_count", 0),
        "stops_count":      stops["count"],
        "stops_seconds":    stops["total_seconds"],
        "history_trips":    history["trip_count"] if history else 0,
        "created_by":       created_by,
        "created_by_email": created_by_email,
    }
    try:
        output_path.with_suffix(".meta.json").write_text(
            json.dumps(meta, indent=2), encoding="utf-8")
    except OSError as exc:
        log.warning("could not write metadata sidecar: %s", exc)


def build_report(*, imei: str, truck_id: str, driver_name: str,
                 begin: datetime, end: datetime,
                 template_path: Path, output_path: Path,
                 include_history: bool = True) -> dict:
    computed = compute_report(
        imei=imei, truck_id=truck_id, driver_name=driver_name,
        begin=begin, end=end, include_history=include_history,
    )
    write_excel(computed, template_path, output_path)
    return computed