"""
report.py — trip detection, event summarisation, and Excel template filling.

Depends on jimi.py for all API access.
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook

from jimi import api_call, fetch_alarms, fetch_alarms_chunked


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

GBC_GEOID = "791a3bb604ef41cfa302e5dc8f2c356c"
GBC_FENCE_NAME_FALLBACK = "TRIP TIME GBC"

SPEED_LIMIT_KPH = 75.0

# Driver-behaviour alert ids (confirmed against live data).
ALERT_ACCELERATION = "41"   # Sudden Acceleration Alert  (no template row)
ALERT_CORNERING    = "43"   # Sharp Turn Right Alert     -> Hard cornering
ALERT_DECELERATION = "48"   # Sudden Deceleration Alert  -> Hard braking
ALERT_OVERSPEED    = "overSpeed"

MIN_STOP_SECONDS   = 5 * 60        # stops >= 5 min count
REST_BREAK_SECONDS = 2 * 60 * 60   # rests >= 2 h count

# Stop-detection tuning. Adjust after comparing against the parking report.
STATIONARY_SPEED_KPH = 1.0     # below this = stationary
MERGE_GAP_SECONDS    = 5 * 60  # merge stationary runs separated by < this
STARTUP_GRACE_S      = 5 * 60  # ignore stops in the first N s of the trip
ENDPARK_GRACE_S      = 30 * 60 # ignore stops in the last N s of the trip

FMT = "%Y-%m-%d %H:%M:%S"


# ---------------------------------------------------------------------------
# Time helpers
# ---------------------------------------------------------------------------

def utc_str(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime(FMT)


def parse_utc(s: str) -> datetime:
    return datetime.strptime(s, FMT).replace(tzinfo=timezone.utc)


def _to_float(x) -> float:
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def _speed_of(p: dict) -> float:
    """
    Track points expose 'gpsSpeed'; alarm payloads expose 'speed'.
    Try gpsSpeed first, then fall back to speed.
    """
    for key in ("gpsSpeed", "speed"):
        v = p.get(key)
        if v is not None:
            try:
                return float(v)
            except (TypeError, ValueError):
                pass
    return 0.0


def history_window(reference: datetime) -> tuple[datetime, datetime]:
    """
    (begin, end) for the previous 3 full calendar months before reference's
    month, clipped to the API's ~90-day retention floor.

    begin inclusive, end exclusive; both UTC-aware.
    """
    first_of_this_month = reference.replace(
        day=1, hour=0, minute=0, second=0, microsecond=0)
    end = first_of_this_month

    # Step back 3 calendar months.
    y, m = first_of_this_month.year, first_of_this_month.month - 3
    while m <= 0:
        m += 12
        y -= 1
    calendar_begin = first_of_this_month.replace(year=y, month=m)

    # Clip to the retention floor — the API won't serve older data.
    retention_floor = reference - timedelta(days=89)
    begin = max(calendar_begin, retention_floor)
    return begin, end


# ---------------------------------------------------------------------------
# Trip detection
# ---------------------------------------------------------------------------

def _is_gbc_alarm(alarm: dict) -> bool:
    if (alarm.get("geoid") or "").strip() == GBC_GEOID:
        return True
    return GBC_FENCE_NAME_FALLBACK in (alarm.get("alarmTypeName") or "")


def find_trips(imei: str, begin: datetime, end: datetime) -> list[dict]:
    alarms = fetch_alarms_chunked(imei, begin, end)
    gbc = [a for a in alarms if _is_gbc_alarm(a)]
    gbc.sort(key=lambda a: a["alertTime"])

    trips: list[dict] = []
    open_start: datetime | None = None
    for a in gbc:
        event = (a.get("alertTypeId") or "").lower()
        t = parse_utc(a["alertTime"])
        if event == "in":
            if open_start is not None:
                print(f"[warn] unclosed trip starting {open_start}; restarting")
            open_start = t
        elif event == "out":
            if open_start is None:
                continue
            trips.append({
                "start": open_start, "end": t,
                "duration_s": (t - open_start).total_seconds(),
            })
            open_start = None
    return trips


# ---------------------------------------------------------------------------
# Track helpers
# ---------------------------------------------------------------------------

def fetch_track(imei: str, begin: datetime, end: datetime) -> list[dict]:
    response = api_call("jimi.device.track.list", {
        "imei":       imei,
        "begin_time": utc_str(begin),
        "end_time":   utc_str(end),
    })
    result = response.get("result") or response.get("data") or []
    if isinstance(result, dict):
        result = result.get("list", [])
    return result or []


# ---------------------------------------------------------------------------
# Stop / rest detection — from track points
# ---------------------------------------------------------------------------

def summarise_stops(track: list[dict],
                    trip_start: datetime | None = None,
                    trip_end: datetime | None = None) -> dict:
    """
    Detect stops from track points.

    A stop = a run of points with speed < STATIONARY_SPEED_KPH that:
      - lasts >= MIN_STOP_SECONDS,
      - does not begin within STARTUP_GRACE_S of trip_start,
      - does not end within ENDPARK_GRACE_S of trip_end.

    Consecutive stationary runs separated by < MERGE_GAP_SECONDS are merged.
    """
    if not track:
        return _empty_stops()

    # 1. Sorted list of (time, is_moving) tuples.
    points = sorted(
        ((parse_utc(p["gpsTime"]), _speed_of(p) >= STATIONARY_SPEED_KPH)
         for p in track if p.get("gpsTime")),
        key=lambda x: x[0],
    )
    if not points:
        return _empty_stops()

    # 2. Build candidate intervals of stationarity.
    intervals: list[list[datetime]] = []
    cur_start: datetime | None = None
    cur_last: datetime | None = None
    for t, moving in points:
        if not moving:
            if cur_start is None:
                cur_start = t
            cur_last = t
        else:
            if cur_start is not None:
                intervals.append([cur_start, cur_last])
                cur_start = cur_last = None
    if cur_start is not None:
        intervals.append([cur_start, cur_last])

    # 3. Merge intervals separated by < MERGE_GAP_SECONDS.
    merged: list[list[datetime]] = []
    for iv in intervals:
        if merged and (iv[0] - merged[-1][1]).total_seconds() < MERGE_GAP_SECONDS:
            merged[-1][1] = iv[1]
        else:
            merged.append(list(iv))
    intervals = merged

    # 4. Apply graces and duration threshold.
    stops: list[float] = []
    for start, end in intervals:
        duration = (end - start).total_seconds()
        if duration < MIN_STOP_SECONDS:
            continue
        if trip_start is not None and \
           (start - trip_start).total_seconds() < STARTUP_GRACE_S:
            continue
        if trip_end is not None and \
           (trip_end - end).total_seconds() < ENDPARK_GRACE_S:
            continue
        stops.append(duration)

    rests = [d for d in stops if d >= REST_BREAK_SECONDS]

    return {
        "count":                 len(stops),
        "total_seconds":         sum(stops),
        "longest_seconds":       max(stops) if stops else 0,
        "rest_break_count":      len(rests),
        "rest_break_seconds":    sum(rests),
        "rest_break_longest_s":  max(rests) if rests else 0,
    }


def _empty_stops() -> dict:
    return {
        "count": 0, "total_seconds": 0, "longest_seconds": 0,
        "rest_break_count": 0, "rest_break_seconds": 0,
        "rest_break_longest_s": 0,
    }


# ---------------------------------------------------------------------------
# Event summarisation
# ---------------------------------------------------------------------------

def summarise_events(alarms: list[dict]) -> dict:
    counts = {
        "hard_braking":     0,
        "hard_cornering":   0,
        "speeding_10":      0,
        "speeding_20":      0,
        "speeding_over_20": 0,
    }
    max_speed = {k: 0.0 for k in counts}
    acceleration_count = 0

    for a in alarms:
        atid = (a.get("alertTypeId") or "").strip()
        name = (a.get("alarmTypeName") or "").lower()
        speed = _to_float(a.get("speed"))

        if atid == ALERT_DECELERATION or "sudden deceleration" in name:
            counts["hard_braking"] += 1
            max_speed["hard_braking"] = max(max_speed["hard_braking"], speed)
        elif atid == ALERT_CORNERING or "sharp turn" in name or "cornering" in name:
            counts["hard_cornering"] += 1
            max_speed["hard_cornering"] = max(max_speed["hard_cornering"], speed)
        elif atid == ALERT_ACCELERATION or "sudden acceleration" in name:
            acceleration_count += 1
        elif atid == ALERT_OVERSPEED or "overspeed" in name:
            over = speed - SPEED_LIMIT_KPH
            if over <= 0:
                continue
            bucket = ("speeding_10" if over < 10 else
                      "speeding_20" if over < 20 else
                      "speeding_over_20")
            counts[bucket] += 1
            max_speed[bucket] = max(max_speed[bucket], speed)

    return {
        "counts": counts,
        "max_speed": max_speed,
        "acceleration_count": acceleration_count,
    }


# ---------------------------------------------------------------------------
# History (previous 3 calendar months, aggregated per trip)
# ---------------------------------------------------------------------------

def build_history(imei: str, reference: datetime) -> dict:
    """
    For every trip whose START falls within the previous 3 full calendar
    months before `reference`'s month (clipped to the API's retention floor),
    fetch its alarms + track, summarise, and aggregate per-trip.
    All averages are "per trip".
    """
    hist_begin, hist_end = history_window(reference)
    print(f"[history] calendar window {utc_str(hist_begin)} → {utc_str(hist_end)}")

    trips = find_trips(imei, hist_begin, hist_end)
    trips = [t for t in trips if hist_begin <= t["start"] < hist_end]
    print(f"[history] {len(trips)} trip(s) starting in window")

    stop_counts: list[float] = []
    stop_secs:   list[float] = []
    rest_secs:   list[float] = []

    behav_events = {k: 0 for k in
                    ("hard_braking", "hard_cornering",
                     "speeding_10", "speeding_20", "speeding_over_20")}
    behav_max   = {k: 0.0 for k in behav_events}
    behav_per_trip_max = {k: [] for k in behav_events}

    for t in trips:
        alarms = fetch_alarms_chunked(imei, t["start"], t["end"])
        try:
            track = fetch_track(imei, t["start"], t["end"])
        except Exception as exc:
            print(f"[history] track fetch failed for {t['start']}: {exc}")
            track = []
        s = summarise_stops(track, t["start"], t["end"])
        e = summarise_events(alarms)

        stop_counts.append(s["count"])
        stop_secs.append(s["total_seconds"])
        rest_secs.append(s["rest_break_seconds"])

        for k in behav_events:
            behav_events[k] += e["counts"][k]
            behav_max[k] = max(behav_max[k], e["max_speed"][k])
            behav_per_trip_max[k].append(e["max_speed"][k])

    def _avg(xs):
        return sum(xs) / len(xs) if xs else 0.0

    return {
        "trip_count": len(trips),
        "window_begin": hist_begin,
        "window_end":   hist_end,
        "stops": {
            "avg_count":    _avg(stop_counts),
            "avg_seconds":  _avg(stop_secs),
            "max_seconds":  max(stop_secs) if stop_secs else 0.0,
            "min_seconds":  min(stop_secs) if stop_secs else 0.0,
        },
        "rests": {
            "avg_seconds": _avg(rest_secs),
            "max_seconds": max(rest_secs) if rest_secs else 0.0,
            "min_seconds": min(rest_secs) if rest_secs else 0.0,
        },
        "behaviours": {
            k: {
                "total_events":  behav_events[k],
                "max_speed":     behav_max[k],
                "avg_max_speed": _avg(behav_per_trip_max[k]),
            }
            for k in behav_events
        },
    }


# ---------------------------------------------------------------------------
# Excel filling
# ---------------------------------------------------------------------------

CELL_REPORT_GENERATED = "D4"
CELL_TRIP_START       = "D5"
CELL_TRIP_END         = "D6"
CELL_TRUCK_ID         = "D7"
CELL_DRIVER_NAME      = "D8"

CELL_REST_BREAK_HOURS = "E12"
CELL_STOPS_COUNT      = "E13"
CELL_STOPS_HOURS      = "E14"

CELL_HARD_BRAKING_COUNT = "E18"
CELL_HARD_BRAKING_SPEED = "F18"
CELL_HARD_CORNER_COUNT  = "E19"
CELL_HARD_CORNER_SPEED  = "F19"
CELL_SPEED_10_COUNT     = "E20"
CELL_SPEED_10_SPEED     = "F20"
CELL_SPEED_20_COUNT     = "E21"
CELL_SPEED_20_SPEED     = "F21"
CELL_SPEED_OVER_COUNT   = "E22"
CELL_SPEED_OVER_SPEED   = "F22"

CELL_HIST_REST_AVG_HOURS   = "H12"
CELL_HIST_REST_LONGEST_H   = "I12"
CELL_HIST_REST_SHORTEST_H  = "J12"

CELL_HIST_STOPS_AVG        = "H13"
CELL_HIST_STOPS_LONGEST    = "I13"
CELL_HIST_STOPS_SHORTEST   = "J13"

CELL_HIST_STOPS_AVG_HOURS     = "H14"
CELL_HIST_STOPS_LONGEST_HOURS = "I14"
CELL_HIST_STOPS_SHORTEST_HOURS = "J14"

CELL_HIST_HB_TOTAL = "H18"; CELL_HIST_HB_MAX = "I18"; CELL_HIST_HB_AVG = "J18"
CELL_HIST_HC_TOTAL = "H19"; CELL_HIST_HC_MAX = "I19"; CELL_HIST_HC_AVG = "J19"
CELL_HIST_S10_TOTAL = "H20"; CELL_HIST_S10_MAX = "I20"; CELL_HIST_S10_AVG = "J20"
CELL_HIST_S20_TOTAL = "H21"; CELL_HIST_S20_MAX = "I21"; CELL_HIST_S20_AVG = "J21"
CELL_HIST_SOV_TOTAL = "H22"; CELL_HIST_SOV_MAX = "I22"; CELL_HIST_SOV_AVG = "J22"


def _safe_write(ws, coord: str, value) -> None:
    for mr in ws.merged_cells.ranges:
        if coord in mr:
            anchor = str(mr).split(":")[0]
            if anchor != coord:
                print(f"[fill_template] {coord} inside merge {mr}; → {anchor}")
                ws[anchor] = value
                return
            break
    ws[coord] = value


def _hours(seconds: float) -> float:
    return round(seconds / 3600, 2)


def fill_template(template_path: Path, output_path: Path, *,
                  truck_id: str, driver_name: str,
                  trip_start: datetime | None,
                  trip_end: datetime | None,
                  stops: dict, events: dict,
                  history: dict | None) -> None:
    wb = load_workbook(template_path)
    ws = wb["Sheet1"]

    _safe_write(ws, CELL_REPORT_GENERATED,
                datetime.now(timezone.utc).strftime(FMT))
    _safe_write(ws, CELL_TRUCK_ID,    truck_id)
    _safe_write(ws, CELL_DRIVER_NAME, driver_name)
    if trip_start:
        _safe_write(ws, CELL_TRIP_START, trip_start.strftime(FMT))
    if trip_end:
        _safe_write(ws, CELL_TRIP_END,   trip_end.strftime(FMT))

    _safe_write(ws, CELL_REST_BREAK_HOURS, _hours(stops["rest_break_seconds"]))
    _safe_write(ws, CELL_STOPS_COUNT,      stops["count"])
    _safe_write(ws, CELL_STOPS_HOURS,      _hours(stops["total_seconds"]))

    c, m = events["counts"], events["max_speed"]
    _safe_write(ws, CELL_HARD_BRAKING_COUNT, c["hard_braking"])
    _safe_write(ws, CELL_HARD_BRAKING_SPEED, round(m["hard_braking"], 1))
    _safe_write(ws, CELL_HARD_CORNER_COUNT,  c["hard_cornering"])
    _safe_write(ws, CELL_HARD_CORNER_SPEED,  round(m["hard_cornering"], 1))
    _safe_write(ws, CELL_SPEED_10_COUNT,     c["speeding_10"])
    _safe_write(ws, CELL_SPEED_10_SPEED,     round(m["speeding_10"], 1))
    _safe_write(ws, CELL_SPEED_20_COUNT,     c["speeding_20"])
    _safe_write(ws, CELL_SPEED_20_SPEED,     round(m["speeding_20"], 1))
    _safe_write(ws, CELL_SPEED_OVER_COUNT,   c["speeding_over_20"])
    _safe_write(ws, CELL_SPEED_OVER_SPEED,   round(m["speeding_over_20"], 1))

    if history and history["trip_count"] > 0:
        h = history
        _safe_write(ws, CELL_HIST_REST_AVG_HOURS,  _hours(h["rests"]["avg_seconds"]))
        _safe_write(ws, CELL_HIST_REST_LONGEST_H,  _hours(h["rests"]["max_seconds"]))
        _safe_write(ws, CELL_HIST_REST_SHORTEST_H, _hours(h["rests"]["min_seconds"]))

        _safe_write(ws, CELL_HIST_STOPS_AVG,      round(h["stops"]["avg_count"], 1))
        _safe_write(ws, CELL_HIST_STOPS_LONGEST,
                    round(h["stops"]["max_seconds"] / 60, 1))
        _safe_write(ws, CELL_HIST_STOPS_SHORTEST,
                    round(h["stops"]["min_seconds"] / 60, 1))
        _safe_write(ws, CELL_HIST_STOPS_AVG_HOURS,      _hours(h["stops"]["avg_seconds"]))
        _safe_write(ws, CELL_HIST_STOPS_LONGEST_HOURS,  _hours(h["stops"]["max_seconds"]))
        _safe_write(ws, CELL_HIST_STOPS_SHORTEST_HOURS, _hours(h["stops"]["min_seconds"]))

        b = h["behaviours"]
        _safe_write(ws, CELL_HIST_HB_TOTAL, b["hard_braking"]["total_events"])
        _safe_write(ws, CELL_HIST_HB_MAX,   round(b["hard_braking"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_HB_AVG,   round(b["hard_braking"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_HC_TOTAL, b["hard_cornering"]["total_events"])
        _safe_write(ws, CELL_HIST_HC_MAX,   round(b["hard_cornering"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_HC_AVG,   round(b["hard_cornering"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_S10_TOTAL, b["speeding_10"]["total_events"])
        _safe_write(ws, CELL_HIST_S10_MAX,   round(b["speeding_10"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_S10_AVG,   round(b["speeding_10"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_S20_TOTAL, b["speeding_20"]["total_events"])
        _safe_write(ws, CELL_HIST_S20_MAX,   round(b["speeding_20"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_S20_AVG,   round(b["speeding_20"]["avg_max_speed"], 1))
        _safe_write(ws, CELL_HIST_SOV_TOTAL, b["speeding_over_20"]["total_events"])
        _safe_write(ws, CELL_HIST_SOV_MAX,   round(b["speeding_over_20"]["max_speed"], 1))
        _safe_write(ws, CELL_HIST_SOV_AVG,   round(b["speeding_over_20"]["avg_max_speed"], 1))

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Top-level pipeline
# ---------------------------------------------------------------------------

def build_report(*, imei: str, truck_id: str, driver_name: str,
                 begin: datetime, end: datetime,
                 template_path: Path, output_path: Path,
                 include_history: bool = True) -> dict:
    print(f"[1/6] Finding trips ...")
    trips = find_trips(imei, begin, end)
    print(f"      {len(trips)} trip(s)")

    if trips:
        trip = trips[0]
        trip_start, trip_end = trip["start"], trip["end"]
    else:
        trip_start = trip_end = None

    print(f"[2/6] Fetching track ...")
    track = fetch_track(imei, trip_start or begin, trip_end or end) if trip_start else []
    print(f"      {len(track)} points")

    print(f"[3/6] Fetching alarms ...")
    alarms = fetch_alarms_chunked(imei, trip_start or begin, trip_end or end) if trip_start else []
    print(f"      {len(alarms)} alarms")

    print(f"[4/6] Summarising ...")
    stops  = summarise_stops(track, trip_start, trip_end)
    events = summarise_events(alarms)
    print(f"      stops >=5m: {stops['count']}  "
          f"total {stops['total_seconds']/3600:.2f}h  "
          f"rests>=2h: {stops['rest_break_count']}")

    history = None
    if include_history and trip_start:
        print(f"[5/6] Building 3-month history ...")
        try:
            history = build_history(imei, trip_start)
        except Exception as exc:
            print(f"      [warn] history failed: {exc}")

    print(f"[6/6] Writing {output_path} ...")
    fill_template(
        template_path, output_path,
        truck_id=truck_id, driver_name=driver_name,
        trip_start=trip_start, trip_end=trip_end,
        stops=stops, events=events, history=history,
    )

    return {
        "trips": trips, "stops": stops, "events": events,
        "history": history,
        "trip_start": trip_start, "trip_end": trip_end,
        "alarm_count": len(alarms), "track_count": len(track),
    }